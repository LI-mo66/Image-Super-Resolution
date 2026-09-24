from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "LFMN复现后创新方案周报_2026-09-24.md"
OUTPUT = ROOT / "LFMN复现后创新方案周报_2026-09-24.docx"
ASSET_DIR = ROOT / "_weekly_report_docx_assets"
FONT_PATH = Path(r"C:\Windows\Fonts\simsun.ttc")
FONT_NAME = "宋体"
BODY_SIZE = Pt(12)
BLACK = RGBColor(0, 0, 0)


def compact_mixed_spacing(text: str) -> str:
    """Remove spaces between Chinese, Latin letters and digits as requested."""
    old = None
    while old != text:
        old = text
        text = re.sub(r"(?<=[\u3400-\u9fff])\s+(?=[A-Za-z0-9])", "", text)
        text = re.sub(r"(?<=[A-Za-z0-9])\s+(?=[\u3400-\u9fff])", "", text)
        text = re.sub(r"(?<=[A-Za-z])\s+(?=[0-9])", "", text)
        text = re.sub(r"(?<=[0-9])\s+(?=[A-Za-z])", "", text)
    return text


def clean_md(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    return compact_mixed_spacing(text.strip())


def set_run_font(run, size=BODY_SIZE, bold=None):
    run.font.name = FONT_NAME
    run.font.size = size
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_repeat_table_header(row):
    repeat_table_header(row)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = trPr.find(qn("w:cantSplit"))
    if cant is None:
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    cant.set(qn("w:val"), "true")


def set_paragraph_format(paragraph, first_line=True, after=Pt(5), before=Pt(0)):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.35
    fmt.space_after = after
    fmt.space_before = before
    if first_line:
        fmt.first_line_indent = Pt(24)
    fmt.widow_control = True


def disable_auto_spacing(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        el = pPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            pPr.append(el)
        el.set(qn("w:val"), "0")


def remove_paragraph_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is not None:
        pPr.remove(pBdr)


def add_text_paragraph(doc, text, style=None, first_line=True, bold_lead=None):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, first_line=first_line)
    text = clean_md(text)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    disable_auto_spacing(p)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(clean_md(text))
    set_run_font(r, Pt(10.5))
    disable_auto_spacing(p)
    return p


def add_picture(doc, path: Path, width=Inches(6.75)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=width)
    return p


def set_table_layout_fixed(table):
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_table(doc, rows: list[list[str]]):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_layout_fixed(table)
    header_fill = "DCE6F1"
    alt_fill = "F5F8FC"
    for i, row in enumerate(rows):
        set_row_cant_split(table.rows[i])
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            text = clean_md(row[j] if j < len(row) else "")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 0 or len(text) < 24 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            set_run_font(run, BODY_SIZE, bold=(i == 0))
            disable_auto_spacing(p)
            if i == 0:
                shade_cell(cell, header_fill)
                p.paragraph_format.keep_with_next = True
            elif i % 2 == 0:
                shade_cell(cell, alt_fill)
    set_repeat_table_header(table.rows[0])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第")
    set_run_font(run, Pt(10.5))
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.extend([fldChar1, instrText, fldChar2])
    run2 = paragraph.add_run("页")
    set_run_font(run2, Pt(10.5))


def setup_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.line_spacing = 1.35
    for name, size, bold in (
        ("Title", Pt(20), True),
        ("Heading 1", Pt(16), True),
        ("Heading 2", Pt(14), True),
        ("Heading 3", Pt(12), True),
        ("Heading 4", Pt(12), True),
    ):
        style = styles[name]
        style.font.name = FONT_NAME
        style.font.size = size
        style.font.bold = bold
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    title_style = styles["Title"]
    style_pPr = title_style._element.get_or_add_pPr()
    style_border = style_pPr.find(qn("w:pBdr"))
    if style_border is not None:
        style_pPr.remove(style_border)


FONT_FILE = str(FONT_PATH if FONT_PATH.exists() else Path(r"C:\Windows\Fonts\simsun.ttc"))


def font(size, bold=False):
    candidate = Path(r"C:\Windows\Fonts\simhei.ttf") if bold else Path(FONT_FILE)
    return ImageFont.truetype(str(candidate if candidate.exists() else FONT_FILE), size=size)


def xy(bounds, x, y):
    width, height = bounds
    return int(x * width), int((1 - y) * height)


def draw_centered_multiline(draw, rect, text, fnt, fill="#000000", spacing=10):
    left, top, right, bottom = rect
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, align="center", spacing=spacing)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((left + right - tw) / 2, (top + bottom - th) / 2), text,
                        font=fnt, fill=fill, align="center", spacing=spacing)


def draw_box(draw, bounds, x, y, w, h, text, fc="#EAF1F8", ec="#365F91", fs=34):
    width, height = bounds
    left, top = xy(bounds, x, y + h)
    right, bottom = xy(bounds, x + w, y)
    draw.rounded_rectangle((left, top, right, bottom), radius=22, fill=fc, outline=ec, width=5)
    draw_centered_multiline(draw, (left + 12, top + 8, right - 12, bottom - 8), text, font(fs), spacing=9)


def arrow(draw, bounds, x1, y1, x2, y2, label=None):
    p1 = xy(bounds, x1, y1)
    p2 = xy(bounds, x2, y2)
    draw.line((p1, p2), fill="#4F4F4F", width=5)
    import math
    angle = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    size = 18
    a = (p2[0] - size * math.cos(angle - 0.55), p2[1] - size * math.sin(angle - 0.55))
    b = (p2[0] - size * math.cos(angle + 0.55), p2[1] - size * math.sin(angle + 0.55))
    draw.polygon((p2, a, b), fill="#4F4F4F")
    if label:
        mx, my = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 - 35
        box = draw.textbbox((0, 0), label, font=font(26))
        draw.text((mx - (box[2] - box[0]) / 2, my), label, font=font(26), fill="#333333")


def canvas(size=(2400, 1000)):
    img = Image.new("RGB", size, "white")
    return img, ImageDraw.Draw(img)


def title(draw, bounds, text):
    width, _ = bounds
    bb = draw.textbbox((0, 0), text, font=font(48, bold=True))
    draw.text(((width - (bb[2] - bb[0])) / 2, 35), text, font=font(48, bold=True), fill="#000000")


def figure_overall(path: Path):
    img, draw = canvas((2400, 940))
    bounds = img.size
    nodes = [
        (0.02, 0.60, 0.15, 0.23, "LFMN复现\n统一数据与评测"),
        (0.22, 0.60, 0.17, 0.23, "预训练权重\n5/10轮同时微调"),
        (0.44, 0.60, 0.17, 0.23, "B0与候选\n从零训练20轮"),
        (0.66, 0.60, 0.15, 0.23, "蒸馏线\nC1/M0/M1"),
        (0.84, 0.60, 0.14, 0.23, "创新确认\n长训与多数据集"),
    ]
    for n in nodes:
        draw_box(draw, bounds, *n)
    for i in range(len(nodes) - 1):
        x1 = nodes[i][0] + nodes[i][2]
        x2 = nodes[i + 1][0]
        arrow(draw, bounds, x1, 0.715, x2, 0.715)
    draw_box(draw, bounds, 0.47, 0.13, 0.26, 0.23, "SRPR结构线\n状态重建→轻量化→低秩门控", fc="#F4F0E6", ec="#8B6B2D")
    arrow(draw, bounds, 0.525, 0.60, 0.565, 0.36, "结构探索")
    arrow(draw, bounds, 0.73, 0.36, 0.89, 0.60, "通过门槛后组合")
    title(draw, bounds, "复现后实验路线与决策门槛")
    img.save(path, dpi=(240, 240))


def figure_structure(path: Path):
    img, draw = canvas((2400, 1080))
    bounds = img.size
    draw_box(draw, bounds, 0.03, 0.58, 0.13, 0.18, "LR输入\n浅层特征")
    draw_box(draw, bounds, 0.20, 0.58, 0.13, 0.18, "阶段1-2")
    draw_box(draw, bounds, 0.37, 0.58, 0.13, 0.18, "阶段3-4")
    draw_box(draw, bounds, 0.54, 0.58, 0.13, 0.18, "阶段5-6")
    draw_box(draw, bounds, 0.71, 0.58, 0.13, 0.18, "阶段7-8")
    draw_box(draw, bounds, 0.87, 0.58, 0.10, 0.18, "重建头\nSR输出")
    xs = [0.16, 0.33, 0.50, 0.67, 0.84, 0.87]
    for i in range(5):
        arrow(draw, bounds, xs[i], 0.67, xs[i + 1], 0.67)
    annotations = [
        (0.18, 0.18, 0.18, 0.19, "P01/RDSM\n阶段反馈与残差需求调制"),
        (0.40, 0.18, 0.17, 0.19, "N3\n中途更新共享先验"),
        (0.60, 0.18, 0.17, 0.19, "N4\n跨阶段差分记忆"),
        (0.79, 0.18, 0.18, 0.19, "N6/N7\n方向交换与交叉轴FFN"),
    ]
    for n in annotations:
        draw_box(draw, bounds, *n, fc="#F8F4EA", ec="#8B6B2D", fs=31)
    arrow(draw, bounds, 0.28, 0.37, 0.43, 0.58)
    arrow(draw, bounds, 0.485, 0.37, 0.485, 0.58)
    arrow(draw, bounds, 0.685, 0.37, 0.63, 0.58)
    arrow(draw, bounds, 0.88, 0.37, 0.78, 0.58)
    title(draw, bounds, "主要结构尝试在LFMN中的作用位置")
    img.save(path, dpi=(240, 240))


def figure_distillation(path: Path):
    img, draw = canvas((2400, 1080))
    bounds = img.size
    draw_box(draw, bounds, 0.04, 0.42, 0.13, 0.19, "同一LR输入")
    draw_box(draw, bounds, 0.29, 0.67, 0.22, 0.18, "学生LFMN\n从零训练")
    draw_box(draw, bounds, 0.29, 0.18, 0.22, 0.18, "冻结教师SwinIR-M\n仅训练期加载")
    draw_box(draw, bounds, 0.65, 0.67, 0.16, 0.18, "学生SR输出")
    draw_box(draw, bounds, 0.65, 0.18, 0.16, 0.18, "教师SR输出")
    draw_box(draw, bounds, 0.85, 0.56, 0.13, 0.25, "C1损失\nL1学生对GT\n+0.1L1学生对教师", fc="#E8F3E8", ec="#397A3B", fs=29)
    arrow(draw, bounds, 0.17, 0.52, 0.29, 0.76)
    arrow(draw, bounds, 0.17, 0.50, 0.29, 0.27)
    arrow(draw, bounds, 0.51, 0.76, 0.65, 0.76)
    arrow(draw, bounds, 0.51, 0.27, 0.65, 0.27)
    arrow(draw, bounds, 0.81, 0.76, 0.85, 0.70)
    arrow(draw, bounds, 0.81, 0.27, 0.85, 0.62)
    note = "部署时仅保留学生LFMN，参数量、FLOPs和推理延迟不增加"
    bb = draw.textbbox((0, 0), note, font=font(32))
    draw.text(((bounds[0] - (bb[2] - bb[0])) / 2, 950), note, font=font(32), fill="#000000")
    title(draw, bounds, "C1普通输出蒸馏训练流程")
    img.save(path, dpi=(240, 240))


def figure_c1_result(path: Path):
    labels = ["C1-B0", "M0-B0", "M1-B0", "M1-C1"]
    values = [0.031656, 0.011599, 0.017645, -0.014011]
    colors = ["#397A3B", "#5B9BD5", "#4472C4", "#C0504D"]
    img, draw = canvas((2000, 1000))
    bounds = img.size
    title(draw, bounds, "蒸馏四组最终结果比较")
    left, right, top, bottom = 230, 1880, 150, 820
    zero_y = int(top + (0.04 / 0.06) * (bottom - top))
    draw.line((left, zero_y, right, zero_y), fill="#555555", width=4)
    for tick_val in (-0.02, -0.01, 0, 0.01, 0.02, 0.03, 0.04):
        y = int(top + ((0.04 - tick_val) / 0.06) * (bottom - top))
        draw.line((left, y, right, y), fill="#D9D9D9", width=2)
        draw.text((65, y - 18), f"{tick_val:+.2f}", font=font(27), fill="#333333")
    bar_w = 210
    xs = [400, 790, 1180, 1570]
    for x, label, val, color in zip(xs, labels, values, colors):
        val_y = int(top + ((0.04 - val) / 0.06) * (bottom - top))
        rect = (x, min(zero_y, val_y), x + bar_w, max(zero_y, val_y))
        draw.rectangle(rect, fill=color)
        bb = draw.textbbox((0, 0), label, font=font(30))
        draw.text((x + (bar_w - (bb[2] - bb[0])) / 2, 850), label, font=font(30), fill="#000000")
        val_text = f"{val:+.6f}"
        vb = draw.textbbox((0, 0), val_text, font=font(28))
        ty = val_y - 45 if val >= 0 else val_y + 12
        draw.text((x + (bar_w - (vb[2] - vb[0])) / 2, ty), val_text, font=font(28), fill="#000000")
    draw.text((20, 90), "epoch40PSNR差值(dB)", font=font(29), fill="#000000")
    img.save(path, dpi=(240, 240))


def figure_srpr(path: Path):
    img, draw = canvas((2400, 1000))
    bounds = img.size
    nodes = [
        (0.02, 0.52, 0.17, 0.23, "N11\nHR持久状态\n末期收益消失"),
        (0.23, 0.52, 0.18, 0.23, "N12\nLR状态+主干回写\n20轮正信号"),
        (0.45, 0.52, 0.16, 0.23, "N13\n统一输出状态\n已实现待训练"),
        (0.65, 0.52, 0.15, 0.23, "N15\n完整共享压缩\n性能下降"),
        (0.84, 0.52, 0.14, 0.23, "当前\n阶段独立\n低秩Gate"),
    ]
    for i, n in enumerate(nodes):
        fc, ec = ("#FDECEC", "#A94442") if i in (0, 3) else (("#E8F3E8", "#397A3B") if i == 4 else ("#EAF1F8", "#365F91"))
        draw_box(draw, bounds, *n, fc=fc, ec=ec, fs=30)
    for i in range(len(nodes) - 1):
        arrow(draw, bounds, nodes[i][0] + nodes[i][2], 0.635, nodes[i + 1][0], 0.635)
    note1 = "保留：state、observation、writeback、q的机制证据"
    note2 = "改进原则：不再跨阶段共享，只压缩每阶段48×99Gate"
    draw.text((260, 700), note1, font=font(29), fill="#000000")
    draw.text((1190, 700), note2, font=font(29), fill="#000000")
    title(draw, bounds, "SRPR结构演进与当前推进方向")
    img.save(path, dpi=(240, 240))


def create_figures():
    ASSET_DIR.mkdir(exist_ok=True)
    paths = {
        "overall": ASSET_DIR / "figure_1_overall.png",
        "structure": ASSET_DIR / "figure_2_structure.png",
        "distill": ASSET_DIR / "figure_3_distill.png",
        "c1": ASSET_DIR / "figure_4_c1_results.png",
        "srpr": ASSET_DIR / "figure_5_srpr.png",
    }
    figure_overall(paths["overall"])
    figure_structure(paths["structure"])
    figure_distillation(paths["distill"])
    figure_c1_result(paths["c1"])
    figure_srpr(paths["srpr"])
    return paths


def parse_markdown(doc: Document, text: str, figs: dict[str, Path]):
    lines = text.splitlines()
    i = 0
    inserted = set()
    in_code = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            if stripped:
                add_text_paragraph(doc, "•" + stripped, first_line=False)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
            table_rows = []
            table_rows.append([x.strip() for x in stripped.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, table_rows)
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(clean_md(stripped[2:]))
            set_run_font(r, Pt(20), True)
            remove_paragraph_border(p)
            disable_auto_spacing(p)
        elif stripped.startswith("## "):
            title = clean_md(stripped[3:])
            p = doc.add_paragraph(title, style="Heading 1")
            for r in p.runs:
                set_run_font(r, Pt(16), True)
            disable_auto_spacing(p)
            if title.startswith("二、第一阶段") and "structure" not in inserted:
                add_picture(doc, figs["structure"])
                add_caption(doc, "图2主要结构尝试在LFMN中的作用位置")
                inserted.add("structure")
            if title.startswith("五、蒸馏实验") and "distill" not in inserted:
                add_picture(doc, figs["distill"])
                add_caption(doc, "图3C1普通输出蒸馏训练流程")
                inserted.add("distill")
            if title.startswith("六、SRPR") and "srpr" not in inserted:
                add_picture(doc, figs["srpr"])
                add_caption(doc, "图5SRPR结构演进与当前推进方向")
                inserted.add("srpr")
        elif stripped.startswith("### "):
            p = doc.add_paragraph(clean_md(stripped[4:]), style="Heading 2")
            for r in p.runs:
                set_run_font(r, Pt(14), True)
            disable_auto_spacing(p)
        elif stripped.startswith("#### "):
            p = doc.add_paragraph(clean_md(stripped[5:]), style="Heading 3")
            for r in p.runs:
                set_run_font(r, Pt(12), True)
            disable_auto_spacing(p)
        elif re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(clean_md(re.sub(r"^[-*]\s+", "", stripped)))
            set_run_font(r)
            disable_auto_spacing(p)
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            content = re.sub(r"^\d+\.\s+", "", stripped)
            r = p.add_run(clean_md(content))
            set_run_font(r)
            disable_auto_spacing(p)
        else:
            p = add_text_paragraph(doc, stripped)
            if stripped.startswith("日期："):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(12)
                if "overall" not in inserted:
                    add_picture(doc, figs["overall"])
                    add_caption(doc, "图1复现后实验路线与决策门槛")
                    inserted.add("overall")
            if stripped.startswith("epoch 40 的近似绝对 SSIM") and "c1" not in inserted:
                add_picture(doc, figs["c1"], width=Inches(6.3))
                add_caption(doc, "图4蒸馏四组epoch40 PSNR差值")
                inserted.add("c1")
        i += 1


def build():
    figs = create_figures()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    setup_styles(doc)
    add_page_number(section.footer.paragraphs[0])
    source_text = SOURCE.read_text(encoding="utf-8")
    parse_markdown(doc, source_text, figs)
    for paragraph in doc.paragraphs:
        disable_auto_spacing(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    disable_auto_spacing(paragraph)
    core = doc.core_properties
    core.title = "LFMN复现后研究周报"
    core.subject = "LFMN结构创新 蒸馏实验与SRPR进展"
    core.author = ""
    core.keywords = "LFMN,超分辨率,蒸馏,SRPR,周报"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
